"""
Data Quality Report
--------------------
Connects the CSV Analyzer and the shared quality scorer into one
professional summary: overview, quality score, issues found, and
recommendations. Exportable to Excel and PDF.

Usage:
    python generate_quality_report.py --input customers.csv --output report.xlsx
    python generate_quality_report.py --input customers.csv --output report.pdf
"""
import argparse
import logging
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from core.io_utils import read_table_safe  # noqa: E402
from core.quality import compute_quality_report  # noqa: E402

logging.basicConfig(level=logging.INFO, format="%(message)s")
log = logging.getLogger(__name__)

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_COLOR = colors.HexColor("#1F4E78")


def build_quality_summary(input_path: Path) -> dict:
    """Load a dataset and compute its quality report (dataset name + all quality_report fields)."""
    if not input_path.exists():
        raise FileNotFoundError(f"File not found: {input_path}")
    df = read_table_safe(input_path)
    if df.empty:
        raise ValueError(f"No data found in {input_path}")
    summary = compute_quality_report(df)
    summary["dataset_name"] = input_path.name
    return summary


def export_quality_report_excel(summary: dict, output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Quality Report"

    ws.cell(row=1, column=1, value="DATA QUALITY REPORT").font = Font(bold=True, size=16)
    ws.cell(row=2, column=1, value=f"Dataset: {summary['dataset_name']}")
    ws.cell(row=3, column=1, value=f"Generated: {datetime.now():%Y-%m-%d %H:%M}")

    ws.cell(row=5, column=1, value="Rows").font = Font(bold=True)
    ws.cell(row=5, column=2, value=summary["rows"])
    ws.cell(row=6, column=1, value="Columns").font = Font(bold=True)
    ws.cell(row=6, column=2, value=summary["columns"])
    ws.cell(row=7, column=1, value="Quality Score").font = Font(bold=True)
    ws.cell(row=7, column=2, value=f"{summary['score']}/100")

    row = 9
    ws.cell(row=row, column=1, value="Issues").font = Font(bold=True)
    ws.cell(row=row, column=1).fill = HEADER_FILL
    ws.cell(row=row, column=1).font = HEADER_FONT
    row += 1
    for issue in (summary["issues"] or ["No issues found"]):
        ws.cell(row=row, column=1, value=f"- {issue}")
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Recommendations").font = Font(bold=True)
    ws.cell(row=row, column=1).fill = HEADER_FILL
    ws.cell(row=row, column=1).font = HEADER_FONT
    row += 1
    for rec in summary["recommendations"]:
        ws.cell(row=row, column=1, value=f"- {rec}")
        row += 1

    ws.column_dimensions["A"].width = 60
    wb.save(output_path)


def export_quality_report_pdf(summary: dict, output_path: Path) -> None:
    doc = SimpleDocTemplate(str(output_path), pagesize=A4,
                             leftMargin=0.7 * inch, rightMargin=0.7 * inch,
                             topMargin=0.6 * inch, bottomMargin=0.6 * inch)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Data Quality Report", styles["Title"]))
    elements.append(Paragraph(f"Dataset: {summary['dataset_name']}", styles["Normal"]))
    elements.append(Paragraph(f"Generated on {datetime.now():%Y-%m-%d %H:%M}", styles["Normal"]))
    elements.append(Spacer(1, 0.2 * inch))

    elements.append(Paragraph(f"<b>Rows:</b> {summary['rows']} &nbsp;&nbsp; <b>Columns:</b> {summary['columns']}", styles["Normal"]))
    elements.append(Paragraph(f"<b>Quality Score:</b> {summary['score']}/100", styles["Heading2"]))
    elements.append(Spacer(1, 0.15 * inch))

    elements.append(Paragraph("Issues", styles["Heading3"]))
    issues = summary["issues"] or ["No issues found"]
    elements.append(ListFlowable([ListItem(Paragraph(i, styles["Normal"])) for i in issues], bulletType="bullet"))
    elements.append(Spacer(1, 0.1 * inch))

    elements.append(Paragraph("Recommendations", styles["Heading3"]))
    elements.append(ListFlowable(
        [ListItem(Paragraph(r, styles["Normal"])) for r in summary["recommendations"]], bulletType="bullet"
    ))

    doc.build(elements)


def main():
    parser = argparse.ArgumentParser(description="Generate a data quality report (Excel or PDF).")
    parser.add_argument("--input", required=True, help="Path to the CSV/Excel file")
    parser.add_argument("--output", required=True, help="Output path (.xlsx or .pdf)")
    args = parser.parse_args()

    try:
        summary = build_quality_summary(Path(args.input))
        output_path = Path(args.output)
        if output_path.suffix.lower() == ".pdf":
            export_quality_report_pdf(summary, output_path)
        else:
            export_quality_report_excel(summary, output_path)
        log.info("Quality score: %d/100", summary["score"])
        log.info("Saved to %s", output_path)
    except (FileNotFoundError, ValueError) as e:
        log.error("Error: %s", e)
    except Exception as e:
        log.error("Unexpected error: %s", e)


if __name__ == "__main__":
    main()
