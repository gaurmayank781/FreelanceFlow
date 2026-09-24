import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from generate_report import generate_report


def _write_csv(path: Path, content: str):
    path.write_text(content)


def test_creates_data_and_summary_sheets(tmp_path):
    _write_csv(tmp_path / "in.csv", "product,units\nA,10\nB,20\n")
    out = tmp_path / "report.xlsx"

    summary = generate_report(tmp_path / "in.csv", out)

    assert out.exists()
    wb = load_workbook(out)
    assert "Data" in wb.sheetnames
    assert "Summary" in wb.sheetnames
    assert summary["rows_written"] == 2


def test_no_summary_sheet_when_no_numeric_columns(tmp_path):
    _write_csv(tmp_path / "in.csv", "product,notes\nA,good\nB,ok\n")
    out = tmp_path / "report.xlsx"

    generate_report(tmp_path / "in.csv", out)

    wb = load_workbook(out)
    assert "Summary" not in wb.sheetnames


def test_title_appears_in_data_sheet(tmp_path):
    _write_csv(tmp_path / "in.csv", "product,units\nA,10\n")
    out = tmp_path / "report.xlsx"

    generate_report(tmp_path / "in.csv", out, title="My Report")

    wb = load_workbook(out)
    ws = wb["Data"]
    assert ws.cell(row=1, column=1).value == "My Report"


def test_summary_stats_are_correct(tmp_path):
    _write_csv(tmp_path / "in.csv", "product,units\nA,10\nB,20\n")
    out = tmp_path / "report.xlsx"

    generate_report(tmp_path / "in.csv", out)

    wb = load_workbook(out)
    ws = wb["Summary"]
    row = [c for c in ws.iter_rows(min_row=2, max_row=2, values_only=True)][0]
    assert row[0] == "units"
    assert row[1] == 2       # count
    assert row[2] == 30      # sum
    assert row[3] == 15      # average


def test_missing_file_raises_error(tmp_path):
    with pytest.raises(FileNotFoundError):
        generate_report(tmp_path / "does_not_exist.csv", tmp_path / "out.xlsx")


def test_empty_csv_raises_error(tmp_path):
    _write_csv(tmp_path / "in.csv", "product,units\n")
    with pytest.raises(ValueError):
        generate_report(tmp_path / "in.csv", tmp_path / "out.xlsx")


def test_dashboard_creates_all_sheets_with_date_and_category(tmp_path):
    from generate_report import generate_dashboard_report

    rows = "\n".join(
        f"2024-01-{d:02d},Widget,{100 + d}" for d in range(1, 11)
    )
    _write_csv(tmp_path / "in.csv", f"date,product,revenue\n{rows}\n")
    out = tmp_path / "dash.xlsx"

    summary = generate_dashboard_report(tmp_path / "in.csv", out, title="Sales Dashboard")

    wb = load_workbook(out)
    assert wb.sheetnames == ["Dashboard", "Chart Data", "Data", "Summary"]
    assert summary["rows_written"] == 10
    assert summary["charts_placed"] >= 1
    assert summary["detected_roles"]["date_col"] == "date"
    assert summary["detected_roles"]["primary_numeric_col"] == "revenue"


def test_dashboard_skips_trend_chart_without_date_column(tmp_path):
    from generate_report import generate_dashboard_report

    _write_csv(tmp_path / "in.csv", "product,revenue\nA,100\nB,200\n")
    out = tmp_path / "dash.xlsx"

    summary = generate_dashboard_report(tmp_path / "in.csv", out)

    assert summary["detected_roles"]["date_col"] is None
    assert any("trend chart" in n.lower() for n in summary["notes"])


def test_dashboard_kpis_include_total_and_records(tmp_path):
    from generate_report import generate_dashboard_report

    _write_csv(tmp_path / "in.csv", "product,revenue\nA,100\nB,300\n")
    out = tmp_path / "dash.xlsx"

    summary = generate_dashboard_report(tmp_path / "in.csv", out)

    kpi_labels = [k for k, _ in summary["kpis"]]
    assert "Total revenue" in kpi_labels
    assert any("Records" in label or "Orders" in label for label in kpi_labels)


def test_dashboard_missing_file_raises(tmp_path):
    from generate_report import generate_dashboard_report
    with pytest.raises(FileNotFoundError):
        generate_dashboard_report(tmp_path / "nope.csv", tmp_path / "out.xlsx")


def test_detect_roles_prefers_money_words():
    from generate_report import _detect_roles
    import pandas as pd

    df = pd.DataFrame({"order_id": [1, 2, 3], "revenue": [10, 20, 30], "qty": [100, 200, 300]})
    roles = _detect_roles(df)
    assert roles["primary_numeric_col"] == "revenue"
    assert roles["records_label"] == "Total Orders"
